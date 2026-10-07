"""
Générateur du Classeur Excel : ANALYSE GLOBALE.xlsx
Tous les biens immobiliers répertoriés à Marrakech avec leurs spécifications complètes.
Chaque quartier dans une nouvelle feuille (page) dédiée + Page Sommaire Index cliquable.
"""

import sqlite3
import os
import sys
import re
import time
import shutil
from datetime import datetime

try:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

DB_PATH = "sourcing_listings.db"
OUTPUT_XLSX_SPACE = "ANALYSE GLOBALE.xlsx"
OUTPUT_XLSX_UNDERSCORE = "ANALYSE_GLOBALE.xlsx"

# Palette Graphique Professionnelle
NAVY_HEADER = "0F172A"      # Slate 900
NAVY_SUBHEADER = "1E293B"   # Slate 800
BLUE_ACCENT = "2563EB"      # Blue 600
BLUE_LIGHT = "DBEAFE"       # Blue 100
BORDER_COLOR = "CBD5E1"     # Slate 300
ZEBRA_FILL = "F8FAFC"       # Slate 50

font_title = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
font_subtitle = Font(name="Calibri", size=10, italic=True, color="94A3B8")
font_section = Font(name="Calibri", size=11, bold=True, color="0F172A")
font_header = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
font_bold = Font(name="Calibri", size=10, bold=True, color="0F172A")
font_regular = Font(name="Calibri", size=10, color="1E293B")
font_link = Font(name="Calibri", size=10, color="2563EB", underline="single")
font_back_link = Font(name="Calibri", size=10, bold=True, color="2563EB", underline="single")

fill_header_navy = PatternFill(start_color=NAVY_HEADER, end_color=NAVY_HEADER, fill_type="solid")
fill_header_blue = PatternFill(start_color=BLUE_ACCENT, end_color=BLUE_ACCENT, fill_type="solid")
fill_zebra = PatternFill(start_color=ZEBRA_FILL, end_color=ZEBRA_FILL, fill_type="solid")
fill_summary_card = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

thin_border = Border(
    left=Side(style="thin", color=BORDER_COLOR),
    right=Side(style="thin", color=BORDER_COLOR),
    top=Side(style="thin", color=BORDER_COLOR),
    bottom=Side(style="thin", color=BORDER_COLOR)
)

COL_WIDTHS_LISTINGS = {
    "A": 10,  # ID
    "B": 38,  # Titre
    "C": 22,  # Lien Web
    "D": 16,  # Type de Bien
    "E": 14,  # Transaction
    "F": 16,  # Prix (DH)
    "G": 14,  # Surface (m²)
    "H": 18,  # Prix au m² (DH/m²)
    "I": 11,  # Chambres
    "J": 13,  # Salles de Bain
    "K": 20,  # Plateforme
    "L": 18,  # Type Vendeur
    "M": 32,  # Équipements / Atouts
    "N": 45,  # Description
}

def sanitize_sheet_name(raw_name, used_names):
    """
    Nettoie et tronque le nom d'un onglet Excel :
    - Caractères interdits : \\ / ? * : [ ] '
    - Longueur max : 31 caractères
    - Détection et résolution des doublons
    """
    s = re.sub(r'[\\/\?\*\:\[\]\']', '-', str(raw_name or 'Quartier')).strip()
    s = re.sub(r'-+', '-', s).strip('- ')
    if not s:
        s = "Quartier"
    s = s[:31]

    candidate = s
    counter = 2
    while candidate.lower() in used_names:
        suffix = f" {counter}"
        candidate = f"{s[:31 - len(suffix)]}{suffix}"
        counter += 1
    
    used_names.add(candidate.lower())
    return candidate

def build_analyse_globale():
    print("=" * 75)
    print("🚀 GÉNÉRATION DU CLASSEUR EXCEL : 'ANALYSE GLOBALE'")
    print("=" * 75)
    t0 = time.time()

    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    cur = conn.cursor()

    total_listings = cur.execute("SELECT COUNT(*) FROM sourcing_listings").fetchone()[0]
    print(f"📦 Total des biens à intégrer : {total_listings:,} biens")

    # Récupérer tous les quartiers groupés avec statistiques
    query_quartiers = """
    SELECT 
        COALESCE(NULLIF(TRIM(quartier), ''), 'Non Spécifié / Autre') as quartier_name,
        COUNT(*) as total_biens,
        ROUND(AVG(CASE WHEN price_mad > 0 THEN price_mad END), 0) as avg_price,
        ROUND(AVG(CASE WHEN surface_m2 > 0 THEN surface_m2 END), 1) as avg_surface,
        ROUND(AVG(CASE WHEN surface_m2 >= 15 AND price_mad > 0 THEN CAST(price_mad AS FLOAT) / surface_m2 END), 0) as avg_pm2,
        COUNT(CASE WHEN house_type = 'Appartement' THEN 1 END) as nb_appart,
        COUNT(CASE WHEN house_type = 'Villa' THEN 1 END) as nb_villa,
        COUNT(CASE WHEN house_type = 'Riad' THEN 1 END) as nb_riad,
        COUNT(CASE WHEN house_type = 'Terrain' THEN 1 END) as nb_terrain,
        COUNT(CASE WHEN house_type IN ('Commerce', 'Bureau', 'Local commercial') THEN 1 END) as nb_commerce,
        COUNT(CASE WHEN house_type NOT IN ('Appartement', 'Villa', 'Riad', 'Terrain', 'Commerce', 'Bureau', 'Local commercial') THEN 1 END) as nb_autre
    FROM sourcing_listings
    GROUP BY quartier_name
    ORDER BY total_biens DESC
    """
    quartier_stats = cur.execute(query_quartiers).fetchall()
    print(f"📍 Nombre total de quartiers identifiés : {len(quartier_stats)}")

    wb = openpyxl.Workbook()
    # Supprimer la première feuille par défaut
    wb.remove(wb.active)

    used_sheet_names = set()

    # =========================================================================
    # 1. FEUILLE MAÎTRESSE : SOMMAIRE GÉNÉRAL & TABLE DES MATIÈRES CLIQUABLE
    # =========================================================================
    summary_title = "📑 SOMMAIRE & RECHERCHE"
    used_sheet_names.add(summary_title.lower())
    ws_index = wb.create_sheet(title=summary_title)
    ws_index.views.sheetView[0].showGridLines = True

    # Bannière Index
    ws_index.merge_cells("A1:M1")
    c1 = ws_index["A1"]
    c1.value = "MARRAKECH IMMOBILIER — ANALYSE GLOBALE PAR QUARTIER"
    c1.font = font_title
    c1.fill = fill_header_navy
    c1.alignment = Alignment(horizontal="center", vertical="center")
    ws_index.row_dimensions[1].height = 36

    ws_index.merge_cells("A2:M2")
    c2 = ws_index["A2"]
    c2.value = f"Inventaire intégral de {total_listings:,} biens — Cliquez sur un quartier pour accéder instantanément à ses biens détaillés"
    c2.font = font_subtitle
    c2.fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    c2.alignment = Alignment(horizontal="center", vertical="center")
    ws_index.row_dimensions[2].height = 20

    # En-têtes du Sommaire
    index_headers = [
        "Rang", "Quartier (Lien Feuille)", "Total Biens", "Prix Moyen (DH)", 
        "Surface Moyenne (m²)", "Prix au m² Moyen (DH/m²)", 
        "Appartements", "Villas", "Riads", "Terrains", "Commerces/Bureaux", "Autres", "Accès Direct"
    ]
    for col_idx, h in enumerate(index_headers, start=1):
        cell = ws_index.cell(row=4, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header_navy
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws_index.row_dimensions[4].height = 28

    # Dictionnaire de correspondance quartier -> nom de feuille
    quartier_to_sheet = {}
    for q_data in quartier_stats:
        raw_name = q_data["quartier_name"]
        sheet_name = sanitize_sheet_name(raw_name, used_sheet_names)
        quartier_to_sheet[raw_name] = sheet_name

    # Remplissage du Sommaire avec Liens Internes Excel
    print("✍️ Construction de la feuille Sommaire...")
    for idx, q_data in enumerate(quartier_stats, start=5):
        raw_name = q_data["quartier_name"]
        target_sheet = quartier_to_sheet[raw_name]
        p_avg = int(q_data["avg_price"]) if q_data["avg_price"] else 0
        s_avg = q_data["avg_surface"] if q_data["avg_surface"] else 0
        pm2_avg = int(q_data["avg_pm2"]) if q_data["avg_pm2"] else 0

        ws_index.cell(row=idx, column=1, value=idx - 4).alignment = Alignment(horizontal="center")
        
        # Lien interne vers l'onglet du quartier
        cell_q = ws_index.cell(row=idx, column=2, value=raw_name)
        cell_q.font = font_link
        cell_q.hyperlink = f"#'{target_sheet}'!A1"
        
        ws_index.cell(row=idx, column=3, value=q_data["total_biens"]).number_format = "#,##0"
        ws_index.cell(row=idx, column=3).alignment = Alignment(horizontal="right")
        
        cell_p = ws_index.cell(row=idx, column=4, value=p_avg if p_avg else "")
        if p_avg: cell_p.number_format = '#,##0" DH"'
        cell_p.alignment = Alignment(horizontal="right")

        cell_s = ws_index.cell(row=idx, column=5, value=s_avg if s_avg else "")
        if s_avg: cell_s.number_format = '0.0" m²"'
        cell_s.alignment = Alignment(horizontal="right")

        cell_pm = ws_index.cell(row=idx, column=6, value=pm2_avg if pm2_avg else "")
        if pm2_avg: cell_pm.number_format = '#,##0" DH/m²"'
        cell_pm.alignment = Alignment(horizontal="right")

        ws_index.cell(row=idx, column=7, value=q_data["nb_appart"]).number_format = "#,##0"
        ws_index.cell(row=idx, column=8, value=q_data["nb_villa"]).number_format = "#,##0"
        ws_index.cell(row=idx, column=9, value=q_data["nb_riad"]).number_format = "#,##0"
        ws_index.cell(row=idx, column=10, value=q_data["nb_terrain"]).number_format = "#,##0"
        ws_index.cell(row=idx, column=11, value=q_data["nb_commerce"]).number_format = "#,##0"
        ws_index.cell(row=idx, column=12, value=q_data["nb_autre"]).number_format = "#,##0"

        # Bouton lien d'accès direct
        cell_action = ws_index.cell(row=idx, column=13, value="👉 Ouvrir l'onglet")
        cell_action.font = font_link
        cell_action.alignment = Alignment(horizontal="center")
        cell_action.hyperlink = f"#'{target_sheet}'!A1"

        for c_i in range(1, 14):
            c = ws_index.cell(row=idx, column=c_i)
            c.border = thin_border
            if not c.font or c.font == font_regular:
                c.font = font_regular

    # Largeurs colonnes sommaire
    index_widths = {
        "A": 8, "B": 28, "C": 13, "D": 18, "E": 16, "F": 19, 
        "G": 14, "H": 12, "I": 12, "J": 12, "K": 16, "L": 10, "M": 18
    }
    for col_l, w in index_widths.items():
        ws_index.column_dimensions[col_l].width = w

    ws_index.auto_filter.ref = f"A4:M{len(quartier_stats) + 4}"

    # =========================================================================
    # 2. FEUILLES PAR QUARTIER : 1 FEUILLE DÉDIÉE PAR QUARTIER
    # =========================================================================
    print("🏘️ Création des feuilles par quartier et intégration des données...")
    
    # Préparation des requêtes pour chaque quartier
    listings_headers = [
        "ID", "Titre de l'Annonce", "Lien Direct", "Type de Bien", "Transaction", 
        "Prix (DH)", "Surface (m²)", "Prix au m² (DH/m²)", "Chambres", "Salles de Bain", 
        "Plateforme", "Type Vendeur", "Équipements & Atouts", "Description"
    ]

    total_sheets_created = 0
    total_listings_inserted = 0

    for q_idx, q_stat in enumerate(quartier_stats, start=1):
        raw_name = q_stat["quartier_name"]
        sheet_name = quartier_to_sheet[raw_name]
        ws = wb.create_sheet(title=sheet_name)
        ws.views.sheetView[0].showGridLines = True
        total_sheets_created += 1

        # Bannière Quartier
        ws.merge_cells("A1:N1")
        cell_t = ws["A1"]
        cell_t.value = f"QUARTIER : {raw_name.upper()} ({q_stat['total_biens']:,} Biens)"
        cell_t.font = font_title
        cell_t.fill = fill_header_navy
        cell_t.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 32

        # Sous-bannière avec Lien de Retour vers le Sommaire + KPIs
        p_avg = int(q_stat["avg_price"]) if q_stat["avg_price"] else 0
        pm2_avg = int(q_stat["avg_pm2"]) if q_stat["avg_pm2"] else 0
        s_avg = q_stat["avg_surface"] if q_stat["avg_surface"] else 0

        ws.cell(row=2, column=1, value="🔙 Retour au Sommaire").font = font_back_link
        ws.cell(row=2, column=1).hyperlink = f"#'{summary_title}'!A1"
        ws.cell(row=2, column=1).alignment = Alignment(horizontal="center", vertical="center")

        ws.merge_cells("B2:N2")
        kpi_text = f"Prix Moyen : {p_avg:,} DH  |  Prix au m² Moyen : {pm2_avg:,} DH/m²  |  Surface Moyenne : {s_avg} m²  |  Total Annonces : {q_stat['total_biens']:,}"
        cell_kpi = ws["B2"]
        cell_kpi.value = kpi_text
        cell_kpi.font = font_subtitle
        cell_kpi.fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        cell_kpi.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[2].height = 20

        # Ligne d'en-tête du tableau
        for col_idx, h in enumerate(listings_headers, start=1):
            cell = ws.cell(row=4, column=col_idx, value=h)
            cell.font = font_header
            cell.fill = fill_header_blue
            cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[4].height = 25

        # Récupération des biens de ce quartier
        if raw_name == 'Non Spécifié / Autre':
            query_items = """
                SELECT id, title, url, house_type, transaction_type, price_mad, surface_m2,
                       bedrooms, bathrooms, platform, seller_type, features, description
                FROM sourcing_listings
                WHERE quartier IS NULL OR TRIM(quartier) = ''
                ORDER BY price_mad DESC
            """
            items = cur.execute(query_items).fetchall()
        else:
            query_items = """
                SELECT id, title, url, house_type, transaction_type, price_mad, surface_m2,
                       bedrooms, bathrooms, platform, seller_type, features, description
                FROM sourcing_listings
                WHERE TRIM(quartier) = ?
                ORDER BY price_mad DESC
            """
            items = cur.execute(query_items, (raw_name,)).fetchall()

        # Remplissage des lignes de biens
        for row_offset, item in enumerate(items, start=5):
            price = item["price_mad"] if item["price_mad"] and item["price_mad"] > 0 else None
            surf = item["surface_m2"] if item["surface_m2"] and item["surface_m2"] > 0 else None
            pm2 = int(price / surf) if (price and surf and surf >= 10) else None

            # ID
            ws.cell(row=row_offset, column=1, value=str(item["id"])).alignment = Alignment(horizontal="center")
            # Titre
            ws.cell(row=row_offset, column=2, value=item["title"] or "Propriété Marrakech")
            
            # Lien Cliquable
            cell_url = ws.cell(row=row_offset, column=3, value="🔗 Voir l'annonce")
            if item["url"]:
                cell_url.hyperlink = item["url"]
                cell_url.font = font_link
            cell_url.alignment = Alignment(horizontal="center")

            # Type de bien
            ws.cell(row=row_offset, column=4, value=item["house_type"] or "Non spécifié")
            # Transaction
            ws.cell(row=row_offset, column=5, value=item["transaction_type"] or "Vente").alignment = Alignment(horizontal="center")

            # Prix
            cell_p = ws.cell(row=row_offset, column=6, value=price)
            if price: cell_p.number_format = '#,##0" DH"'
            cell_p.alignment = Alignment(horizontal="right")

            # Surface
            cell_s = ws.cell(row=row_offset, column=7, value=surf)
            if surf: cell_s.number_format = '#,##0" m²"'
            cell_s.alignment = Alignment(horizontal="right")

            # Prix au m²
            cell_pm2 = ws.cell(row=row_offset, column=8, value=pm2)
            if pm2: cell_pm2.number_format = '#,##0" DH/m²"'
            cell_pm2.alignment = Alignment(horizontal="right")

            # Chambres & Salles de bain
            ws.cell(row=row_offset, column=9, value=item["bedrooms"] if item["bedrooms"] else "").alignment = Alignment(horizontal="center")
            ws.cell(row=row_offset, column=10, value=item["bathrooms"] if item["bathrooms"] else "").alignment = Alignment(horizontal="center")

            # Plateforme
            ws.cell(row=row_offset, column=11, value=item["platform"] or "")
            # Vendeur
            ws.cell(row=row_offset, column=12, value=item["seller_type"] or "")
            # Équipements
            ws.cell(row=row_offset, column=13, value=(item["features"] or "")[:120])
            # Description
            clean_desc = (item["description"] or "").replace("\n", " ").strip()[:180]
            ws.cell(row=row_offset, column=14, value=clean_desc)

            # Bordures & Police standard
            for c_i in range(1, 15):
                c = ws.cell(row=row_offset, column=c_i)
                c.border = thin_border
                if c_i != 3:  # ne pas écraser le style du lien
                    c.font = font_regular

            total_listings_inserted += 1

        # Activer les filtres automatiques sur le tableau du quartier
        last_row = len(items) + 4
        ws.auto_filter.ref = f"A4:N{max(last_row, 5)}"

        # Largeurs de colonnes
        for col_l, w in COL_WIDTHS_LISTINGS.items():
            ws.column_dimensions[col_l].width = w

        if q_idx % 25 == 0 or q_idx == len(quartier_stats):
            print(f"  ⚡ [{q_idx}/{len(quartier_stats)}] Quartiers traités ({total_listings_inserted:,} biens insérés)...")

    # Sauvegarde des fichiers
    print(f"\n💾 Enregistrement du fichier '{OUTPUT_XLSX_SPACE}'...")
    wb.save(OUTPUT_XLSX_SPACE)
    
    # Copie également en version sans espace pour compatibilité web URL
    shutil.copyfile(OUTPUT_XLSX_SPACE, OUTPUT_XLSX_UNDERSCORE)

    dur = time.time() - t0
    size_mb = os.path.getsize(OUTPUT_XLSX_SPACE) / (1024 * 1024)
    print("=" * 75)
    print(f"✅ SUCCÈS COMPLET !")
    print(f"📄 Fichier généré : '{OUTPUT_XLSX_SPACE}' & '{OUTPUT_XLSX_UNDERSCORE}'")
    print(f"📑 Feuilles créées : {total_sheets_created} feuilles de quartiers + 1 Sommaire Interactif")
    print(f"🏡 Biens répertoriés : {total_listings_inserted:,} annonces avec spécifications complètes")
    print(f"🔗 Tous les liens hypertextes inclus et fonctionnels")
    print(f"⏱️ Durée d'exécution : {dur:.1f}s | Taille fichier : {size_mb:.2f} MB")
    print("=" * 75)

if __name__ == "__main__":
    build_analyse_globale()
